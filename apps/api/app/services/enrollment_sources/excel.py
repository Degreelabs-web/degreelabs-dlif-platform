from pathlib import Path

from openpyxl import load_workbook

from app.services.enrollment_sources.base import WorkbookRows, normalize_header


class LocalExcelEnrollmentSource:
    source_type = "local"

    def __init__(
        self,
        source: str,
        student_sheet: str,
        mentor_sheet: str,
        *,
        include_students: bool = True,
        include_mentors: bool = True,
        allow_missing_sheets: bool = False,
        mentor_category: str | None = None,
    ) -> None:
        self.path = Path(source).expanduser()
        self.student_sheet = student_sheet
        self.mentor_sheet = mentor_sheet
        self.include_students = include_students
        self.include_mentors = include_mentors
        self.allow_missing_sheets = allow_missing_sheets
        self.mentor_category = mentor_category

    def is_configured(self) -> bool:
        return bool(str(self.path)) and self.path.is_file()

    def read(self) -> WorkbookRows:
        if not self.is_configured():
            raise FileNotFoundError(
                "The configured enrollment workbook could not be found."
            )

        workbook = load_workbook(
            filename=self.path,
            read_only=True,
            data_only=True,
        )
        try:
            students = (
                self._read_student_sheet(workbook)
                if self.include_students
                else []
            )
            mentors = (
                self._read_sheet(
                    workbook,
                    self.mentor_sheet,
                    aliases=("Form responses 1",),
                )
                if self.include_mentors
                else []
            )
            if self.mentor_category:
                mentors = [
                    {**row, "mentor_category": self.mentor_category}
                    for row in mentors
                ]
            return WorkbookRows(students=students, mentors=mentors)
        finally:
            workbook.close()

    def _read_student_sheet(self, workbook) -> list[dict[str, object]]:
        aliases = ("Student Roster", "Student Details", "Student Data")
        configured_names = (self.student_sheet, *aliases)
        if any(name in workbook.sheetnames for name in configured_names):
            return self._read_sheet(
                workbook,
                self.student_sheet,
                aliases=aliases,
            )

        candidates = [
            worksheet.title
            for worksheet in workbook.worksheets
            if self._looks_like_student_roster(worksheet)
        ]
        if len(candidates) == 1:
            return self._read_sheet(workbook, candidates[0])
        if self.allow_missing_sheets and not candidates:
            return []
        if len(candidates) > 1:
            raise ValueError(
                "Workbook has multiple possible student roster sheets. "
                "Set ENROLLMENT_STUDENT_SHEET to the intended sheet name."
            )
        raise ValueError(
            "Workbook must contain a student roster with Student ID, Student Name, "
            "and Email columns."
        )

    @staticmethod
    def _looks_like_student_roster(worksheet) -> bool:
        rows = worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
        header_row = next(rows, ())
        headers = {normalize_header(value) for value in header_row}
        has_student_id = bool(headers & {"student_id", "roll_number", "roll_no"})
        has_name = bool(headers & {"student_name", "full_name", "name"})
        return has_student_id and has_name and "email" in headers

    def _read_sheet(
        self,
        workbook,
        sheet_name: str,
        *,
        aliases: tuple[str, ...] = (),
    ) -> list[dict[str, object]]:
        resolved_sheet_name = next(
            (
                candidate
                for candidate in (sheet_name, *aliases)
                if candidate in workbook.sheetnames
            ),
            None,
        )
        if resolved_sheet_name is None:
            if self.allow_missing_sheets:
                return []
            raise ValueError(f"Workbook sheet '{sheet_name}' was not found.")

        worksheet = workbook[resolved_sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return []

        headers = [normalize_header(value) for value in header_row]
        if not any(headers):
            return []

        result: list[dict[str, object]] = []
        for row_number, values in enumerate(rows, start=2):
            record = {
                header: value
                for header, value in zip(headers, values, strict=False)
                if header
            }
            if any(value not in (None, "") for value in record.values()):
                record["_row_number"] = row_number
                result.append(record)
        return result
