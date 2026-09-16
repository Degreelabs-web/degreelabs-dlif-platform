from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.enrollment_sources.base import normalize_header


class InstitutionWorkbookError(ValueError):
    pass


def parse_institution_workbook(contents: bytes) -> list[dict[str, str | None]]:
    """Read a simple institution workbook without relying on tab names."""
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception as exc:
        raise InstitutionWorkbookError("Upload a valid .xlsx workbook.") from exc

    for worksheet in workbook.worksheets:
        values = worksheet.iter_rows(values_only=True)
        header_row = next(values, None)
        if not header_row:
            continue

        headers = [normalize_header(str(value or "")) for value in header_row]
        name_index = _header_index(headers, "institution_name", "college_name", "name")
        if name_index is None:
            continue

        address_index = _header_index(headers, "address", "college_address", "location")
        code_index = _header_index(headers, "institution_code", "college_code", "code")
        status_index = _header_index(headers, "status")
        rows: list[dict[str, str | None]] = []

        for row_number, values_row in enumerate(values, start=2):
            if not any(value not in (None, "") for value in values_row):
                continue
            rows.append(
                {
                    "row_number": str(row_number),
                    "name": _cell(values_row, name_index),
                    "address": _cell(values_row, address_index),
                    "code": _cell(values_row, code_index),
                    "status": _cell(values_row, status_index),
                }
            )

        if not rows:
            raise InstitutionWorkbookError(
                f"The '{worksheet.title}' sheet has headers but no institution rows."
            )
        return rows

    raise InstitutionWorkbookError(
        "Workbook needs an Institution Name, College Name, or Name column."
    )


def _header_index(headers: list[str], *aliases: str) -> int | None:
    for alias in aliases:
        if alias in headers:
            return headers.index(alias)
    return None


def _cell(row: tuple[Any, ...], index: int | None) -> str | None:
    if index is None or index >= len(row) or row[index] is None:
        return None
    value = row[index]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None
