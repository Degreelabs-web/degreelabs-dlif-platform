from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from uuid import NAMESPACE_URL, uuid4, uuid5

import pytest
from openpyxl import Workbook
from sqlalchemy import JSON, create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import Settings
from app.db.models.enrollment_sync_run import EnrollmentSyncRun
from app.db.models.institution import Institution
from app.db.models.mentor import Mentor
from app.db.models.student_profile import StudentProfile
from app.db.models.team_mentor_assignment import TeamMentorAssignment
from app.db.models.user import User
from app.services.enrollment_sources.base import WorkbookRows
from app.services.enrollment_sources.excel import LocalExcelEnrollmentSource
from app.services.enrollment_sources.factory import build_enrollment_source
from app.services.enrollment_sources.google_sheets import GoogleSheetsMentorSource
from app.services.enrollment_sync import EnrollmentSyncService
from app.services.enrollment_upload import (
    EnrollmentUploadError,
    EnrollmentWorkbookUploadService,
)
from app.services.mentor import MentorService
from app.schemas.mentor import MentorUpdate


@dataclass
class MutableSource:
    students: list[dict[str, object]]
    mentors: list[dict[str, object]]
    source_type: str = "test"

    def is_configured(self) -> bool:
        return True

    def read(self) -> WorkbookRows:
        return WorkbookRows(students=self.students, mentors=self.mentors)


class FakeSupabaseAdmin:
    def __init__(self) -> None:
        self.users: dict[str, dict] = {}
        self.deleted_ids: list[str] = []

    def get_user_by_email(self, email: str) -> dict | None:
        return self.users.get(email.strip().lower())

    def invite_user(
        self,
        *,
        email: str,
        full_name: str,
        redirect_to: str | None = None,
    ) -> tuple[dict, bool]:
        del redirect_to
        normalized = email.strip().lower()
        existing = self.users.get(normalized)
        if existing:
            return existing, False
        user = {
            "id": str(uuid5(NAMESPACE_URL, normalized)),
            "email": normalized,
            "user_metadata": {"full_name": full_name},
        }
        self.users[normalized] = user
        return user, True

    def update_user(
        self,
        user_id: str,
        *,
        password: str | None = None,
        full_name: str | None = None,
    ) -> dict:
        del password
        user = next(item for item in self.users.values() if item["id"] == user_id)
        if full_name:
            user["user_metadata"] = {"full_name": full_name}
        return user

    def delete_user(self, user_id: str) -> None:
        self.deleted_ids.append(user_id)
        for email, user in list(self.users.items()):
            if user["id"] == user_id:
                del self.users[email]


class FakeGoogleSheetsResponse:
    status_code = 200

    def __init__(self, values: list[list[object]]) -> None:
        self.values = values

    def json(self) -> dict[str, list[list[object]]]:
        return {"values": self.values}


class FakeGoogleSheetsClient:
    def __init__(self, values: list[list[object]]) -> None:
        self.values = values
        self.last_headers: dict[str, str] | None = None

    def get(self, _url: str, **kwargs) -> FakeGoogleSheetsResponse:
        self.last_headers = kwargs.get("headers")
        return FakeGoogleSheetsResponse(self.values)


@pytest.fixture()
def db() -> Session:
    array_columns = ("expertise", "industries", "support_preferences")
    original_array_types = {
        column: Mentor.__table__.c[column].type for column in array_columns
    }
    for column in array_columns:
        Mentor.__table__.c[column].type = JSON()
    engine = create_engine("sqlite+pysqlite:///:memory:")
    for table in (
        User.__table__,
        Institution.__table__,
        StudentProfile.__table__,
        Mentor.__table__,
        TeamMentorAssignment.__table__,
        EnrollmentSyncRun.__table__,
    ):
        table.create(engine)

    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()
        for column, column_type in original_array_types.items():
            Mentor.__table__.c[column].type = column_type


def build_service(
    db: Session,
    source: MutableSource,
    supabase: FakeSupabaseAdmin,
    **config_overrides: object,
) -> EnrollmentSyncService:
    return EnrollmentSyncService(
        db,
        source=source,
        supabase=supabase,  # type: ignore[arg-type]
        config=Settings(
            _env_file=None,
            database_url="sqlite://",
            supabase_url="https://example.supabase.co",
            supabase_service_role_key="test-service-role-key-that-is-long-enough",
            **config_overrides,
        ),
    )


def add_institution(db: Session) -> Institution:
    institution = Institution(
        id=uuid4(),
        name="Apex Institute of Technology",
        code="AIT",
        status="active",
    )
    db.add(institution)
    db.commit()
    return institution


def student_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "_row_number": 2,
        "full_name": "Aarav Sharma",
        "email": " AARAV@EXAMPLE.EDU ",
        "roll_number": "AIT-001",
        "institution_code": "AIT",
        "course": "B.Tech",
        "branch": "Computer Science",
        "graduation_year": 2027,
        "status": "Enrolled",
    }
    row.update(overrides)
    return row


def mentor_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "_row_number": 2,
        "full_name": "Priya Nair",
        "email": "PRIYA@EXAMPLE.COM",
        "company": "DeepMind",
        "designation": "Research Scientist",
        "expertise": "AI; Reinforcement Learning",
        "status": "active",
    }
    row.update(overrides)
    return row


def test_new_student_is_created_and_second_sync_is_idempotent(db: Session) -> None:
    add_institution(db)
    source = MutableSource(students=[student_row()], mentors=[])
    supabase = FakeSupabaseAdmin()
    service = build_service(db, source, supabase)

    first = service.run(use_lock=False)
    second = service.run(use_lock=False)

    assert first.students_created == 1
    assert second.students_created == 0
    assert second.students_updated == 0
    assert db.scalar(select(func.count(User.id))) == 1
    assert db.scalar(select(func.count(StudentProfile.id))) == 1


def test_changed_student_is_updated_without_duplication(db: Session) -> None:
    add_institution(db)
    source = MutableSource(students=[student_row()], mentors=[])
    supabase = FakeSupabaseAdmin()
    service = build_service(db, source, supabase)
    service.run(use_lock=False)

    source.students = [student_row(course="B.Tech AI", phone="1234567890")]
    result = service.run(use_lock=False)
    profile = db.scalar(select(StudentProfile))

    assert result.students_updated == 1
    assert profile is not None
    assert profile.course == "B.Tech AI"
    assert profile.phone == "1234567890"
    assert db.scalar(select(func.count(User.id))) == 1


def test_invalid_student_is_skipped_and_reported(db: Session) -> None:
    add_institution(db)
    source = MutableSource(
        students=[student_row(email="not-an-email")],
        mentors=[],
    )
    result = build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)

    assert result.status == "partial"
    assert result.rows_skipped == 1
    assert result.validation_errors[0]["sheet"] == "student"
    assert db.scalar(select(func.count(User.id))) == 0


def test_new_mentor_is_created_and_second_sync_is_idempotent(db: Session) -> None:
    source = MutableSource(students=[], mentors=[mentor_row()])
    supabase = FakeSupabaseAdmin()
    service = build_service(db, source, supabase)

    first = service.run(use_lock=False)
    second = service.run(use_lock=False)

    assert first.mentors_created == 1
    assert second.mentors_created == 0
    assert second.mentors_updated == 0
    assert db.scalar(select(func.count(User.id))) == 1
    assert db.scalar(select(func.count(Mentor.id))) == 1


def test_changed_mentor_is_updated_without_duplication(db: Session) -> None:
    source = MutableSource(students=[], mentors=[mentor_row()])
    supabase = FakeSupabaseAdmin()
    service = build_service(db, source, supabase)
    service.run(use_lock=False)

    source.mentors = [mentor_row(designation="Principal Scientist")]
    result = service.run(use_lock=False)
    mentor = db.scalar(select(Mentor))

    assert result.mentors_updated == 1
    assert mentor is not None
    assert mentor.designation == "Principal Scientist"
    assert db.scalar(select(func.count(Mentor.id))) == 1


def test_mentor_form_response_fields_are_mapped_and_persisted(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[
            {
                "_row_number": 2,
                "email_address": "mentor@example.com",
                "full_name": "Ananya Rao",
                "current_role": "Vice President, Product",
                "organisation": "Example Labs",
                "city": "Bengaluru",
                "country": "India",
                "mobile_number": "+91 9876543210",
                "linkedin_or_professional_profile_link": "https://linkedin.com/in/ananya",
                "professional_headline": "Product leader and startup mentor",
                "short_professional_bio": "Builds inclusive technology products.",
                "professional_headshot": "https://example.com/ananya.jpg",
                "areas_of_expertise": "Product Strategy; Leadership; product strategy",
                "industries": "Technology, Education, technology",
                "years_of_professional_experience": "15+ years",
                "how_would_you_like_to_support_dlif_fellows": "Career guidance; Product reviews",
                "complete_this_sentence_as_a_mentor_i_help_fellows": "turn ideas into practical plans.",
            }
        ],
    )

    result = build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)
    mentor = db.scalar(select(Mentor))

    assert result.mentors_created == 1
    assert mentor is not None
    assert mentor.designation == "Vice President, Product"
    assert mentor.company_name == "Example Labs"
    assert mentor.city == "Bengaluru"
    assert mentor.country == "India"
    assert mentor.location == "Bengaluru, India"
    assert mentor.phone == "+91 9876543210"
    assert mentor.professional_headline == "Product leader and startup mentor"
    assert mentor.bio == "Builds inclusive technology products."
    assert mentor.expertise == ["Product Strategy", "Leadership"]
    assert mentor.industries == ["Technology", "Education"]
    assert mentor.years_of_experience == 15
    assert mentor.support_preferences == ["Career guidance", "Product reviews"]
    assert mentor.mentor_statement == "turn ideas into practical plans."


def test_mentor_form_response_without_optional_headshot_is_created(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[
            {
                "_row_number": 2,
                "email_address": "mentor@example.com",
                "full_name": "Ananya Rao",
                "current_role": "Product Lead",
                "organisation": "Example Labs",
            }
        ],
    )

    result = build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)
    mentor = db.scalar(select(Mentor))

    assert result.mentors_created == 1
    assert mentor is not None
    assert mentor.headshot_url is None


def test_same_mentor_email_submitted_twice_updates_existing_profile(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[
            mentor_row(
                email="mentor@example.com",
                company="Example Labs",
                designation="Product Lead",
            )
        ],
    )
    service = build_service(db, source, FakeSupabaseAdmin())
    service.run(use_lock=False)

    source.mentors = [
        mentor_row(
            email=" MENTOR@example.com ",
            company="Updated Labs",
            designation="Director",
        )
    ]
    result = service.run(use_lock=False)
    mentor = db.scalar(select(Mentor))

    assert result.mentors_updated == 1
    assert db.scalar(select(func.count(Mentor.id))) == 1
    assert mentor is not None
    assert mentor.company_name == "Updated Labs"
    assert mentor.designation == "Director"


def test_mentor_approval_mode_creates_pending_and_admin_can_approve(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[mentor_row(status=None)],
    )
    service = build_service(
        db,
        source,
        FakeSupabaseAdmin(),
        mentor_onboarding_requires_approval=True,
    )

    service.run(use_lock=False)
    mentor = db.scalar(select(Mentor))
    user = db.scalar(select(User))

    assert mentor is not None
    assert user is not None
    assert mentor.status == "pending"
    assert user.status == "pending"

    updated = MentorService(db).update(
        mentor.id,
        MentorUpdate(status="active"),
    )

    assert updated is not None
    assert updated.status == "active"
    assert user.status == "active"


def test_negative_mentor_experience_is_rejected(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[mentor_row(years_of_professional_experience="-5 years")],
    )

    result = build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)

    assert result.rows_skipped == 1
    assert db.scalar(select(func.count(Mentor.id))) == 0


def test_enrolled_mentors_response_returns_onboarding_fields_and_filters(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[
            {
                "_row_number": 2,
                "email_address": "rahul@example.com",
                "full_name": "Rahul Sharma",
                "current_role": "Engineering Manager",
                "organisation": "Microsoft",
                "city": "Bengaluru",
                "country": "India",
                "mobile_number": "+91 9876543210",
                "linkedin_or_professional_profile_link": "https://linkedin.com/in/rahul",
                "professional_headline": "Engineering Leader | AI Mentor",
                "short_professional_bio": "Builds AI platforms.",
                "professional_headshot": "https://example.com/rahul.jpg",
                "areas_of_expertise": "AI, Python",
                "industries": "Technology",
                "years_of_professional_experience": "10",
                "complete_this_sentence_as_a_mentor_i_help_fellows": "build real-world engineering skills.",
                "status": "active",
            }
        ],
    )
    build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)

    mentors = MentorService(db).get_all(
        status="active",
        organisation="Microsoft",
        industry="Tech",
        expertise="Pyth",
        country="India",
        search="Rahul",
    )

    assert len(mentors) == 1
    response = mentors[0]
    assert response.email == "rahul@example.com"
    assert response.full_name == "Rahul Sharma"
    assert response.current_role == "Engineering Manager"
    assert response.organisation == "Microsoft"
    assert response.city == "Bengaluru"
    assert response.country == "India"
    assert response.professional_headshot_url == "https://example.com/rahul.jpg"
    assert response.mentoring_statement == "build real-world engineering skills."
    detail = MentorService(db).get_by_id(response.id)
    assert detail is not None
    assert detail.phone == "+91 9876543210"
    assert detail.bio == "Builds AI platforms."
    assert detail.expertise == ["AI", "Python"]
    assert detail.industries == ["Technology"]


def test_invalid_mentor_is_skipped_and_reported(db: Session) -> None:
    source = MutableSource(
        students=[],
        mentors=[mentor_row(full_name="")],
    )
    result = build_service(db, source, FakeSupabaseAdmin()).run(use_lock=False)

    assert result.status == "partial"
    assert result.rows_skipped == 1
    assert result.validation_errors[0]["sheet"] == "mentor"
    assert db.scalar(select(func.count(Mentor.id))) == 0


def test_local_excel_source_reads_configured_sheets(tmp_path) -> None:
    workbook_path = tmp_path / "enrollments.xlsx"
    workbook = Workbook()
    student_sheet = workbook.active
    student_sheet.title = "Student Roster"
    student_sheet.append(["Full Name", "Email", "Roll Number"])
    student_sheet.append(["Aarav Sharma", "aarav@example.edu", "AIT-001"])
    mentor_sheet = workbook.create_sheet("Mentor Roster")
    mentor_sheet.append([
        "Full Name",
        "Email",
        "Company",
        "City, Country",
        "Complete this sentence: “As a mentor, I help fellows…”",
    ])
    mentor_sheet.append([
        "Priya Nair",
        "priya@example.com",
        "DeepMind",
        "Bengaluru, India",
        "build confidence.",
    ])
    workbook.save(workbook_path)

    rows = LocalExcelEnrollmentSource(
        str(workbook_path),
        "Student Roster",
        "Mentor Roster",
    ).read()

    assert rows.students[0]["full_name"] == "Aarav Sharma"
    assert rows.students[0]["_row_number"] == 2
    assert rows.mentors[0]["company"] == "DeepMind"
    assert rows.mentors[0]["city_country"] == "Bengaluru, India"
    assert (
        rows.mentors[0]["complete_this_sentence_as_a_mentor_i_help_fellows"]
        == "build confidence."
    )


def test_google_sheets_source_reads_form_response_headers() -> None:
    client = FakeGoogleSheetsClient(
        [
            [
                "Timestamp",
                "Email address",
                "Full name",
                "City, Country",
                "How would you like to support DLIF fellows?",
            ],
            [
                "11/09/2026 09:30:00",
                "mentor@example.com",
                "Ananya Rao",
                "Bengaluru, India",
                "Career guidance, Product reviews",
            ],
        ]
    )
    source = GoogleSheetsMentorSource(
        spreadsheet_id="sheet-id",
        mentor_range="Form responses 1!A:AZ",
        service_account_file="",
        access_token_provider=lambda: "test-token",
        http_client=client,  # type: ignore[arg-type]
    )

    rows = source.read()

    assert rows.students == []
    assert rows.mentors[0]["email_address"] == "mentor@example.com"
    assert rows.mentors[0]["city_country"] == "Bengaluru, India"
    assert (
        rows.mentors[0]["how_would_you_like_to_support_dlif_fellows"]
        == "Career guidance, Product reviews"
    )
    assert rows.mentors[0]["_row_number"] == 2
    assert client.last_headers == {"Authorization": "Bearer test-token"}


def test_factory_builds_google_sheets_source() -> None:
    source = build_enrollment_source(
        Settings(
            _env_file=None,
            enrollment_excel_source_type="google_sheets",
            enrollment_google_spreadsheet_id="sheet-id",
            enrollment_google_mentor_range="Form responses 1!A:AZ",
            enrollment_google_service_account_file="credentials.json",
        )
    )

    assert isinstance(source, GoogleSheetsMentorSource)


def workbook_bytes() -> BytesIO:
    output = BytesIO()
    workbook = Workbook()
    students = workbook.active
    students.title = "Students"
    students.append(["full_name", "email", "roll_number", "institution_code"])
    students.append(["Aarav Sharma", "aarav@example.edu", "AIT-001", "AIT"])
    mentors = workbook.create_sheet("Mentors")
    mentors.append(["full_name", "email", "company"])
    mentors.append(["Priya Nair", "priya@example.com", "DeepMind"])
    workbook.save(output)
    output.seek(0)
    return output


def test_uploaded_workbook_atomically_becomes_master_source(tmp_path) -> None:
    destination = tmp_path / "master.xlsx"
    service = EnrollmentWorkbookUploadService(
        Settings(
            _env_file=None,
            enrollment_excel_source=str(destination),
        )
    )

    saved_path = service.save("enrollments.xlsx", workbook_bytes())

    assert saved_path == destination
    assert destination.is_file()
    assert LocalExcelEnrollmentSource(str(destination), "Students", "Mentors").read().students


def test_upload_rejects_non_xlsx_files(tmp_path) -> None:
    service = EnrollmentWorkbookUploadService(
        Settings(
            _env_file=None,
            enrollment_excel_source=str(tmp_path / "master.xlsx"),
        )
    )

    with pytest.raises(EnrollmentUploadError, match="Only .xlsx"):
        service.save("students.csv", BytesIO(b"not a workbook"))
